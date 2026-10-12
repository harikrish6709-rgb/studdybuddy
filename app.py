import hashlib
import os
import re
import sys
import traceback

import streamlit as st

import database as db
from export import safe_name
from files import ALLOWED_TYPES, make_attachment
from llm import generate_text, stream_reply
from prompts import LEVELS, QUICK_ACTIONS, build_system_prompt
from search import search_web


# ---------------- Make the content for PPT / PDF / Word (plain outline, no JSON) ----------------
OUTLINE_SYSTEM = ("You write clear, accurate study material for students. "
                  "Follow the student's request EXACTLY, including every extra instruction. "
                  "Follow the requested format exactly. No introduction and no closing remarks.")

OUTLINE_FORMAT = (
    "Reply in EXACTLY this plain-text format and nothing else:\n"
    "# Title of the document\n"
    "## Section heading\n"
    "- one short sentence\n"
    "- another short sentence\n"
    "```java\n"
    "short code example (only where useful)\n"
    "```\n"
    "Put 3 to 6 points under every section heading. Each point is ONE short sentence on ONE line. "
    "If the topic is about programming, or the student asks for code or programming examples, "
    "add a short, correct code example in a ``` code block under the relevant sections "
    "(maximum 14 lines, maximum 55 characters per line, in the language the student asked for). "
    "If the student asks for a diagram, draw a simple text diagram inside a ``` block. "
    "Do not use bold, numbering or tables."
)


def _clean_line(t: str) -> str:
    return t.replace("**", "").replace("`", "").strip()


def parse_outline(text: str):
    """Read '# title', '## heading', '- point' and ``` code blocks. Very forgiving."""
    title, sections = None, []
    in_code, code = False, []

    def close_code():
        block = "\n".join(code).strip("\n")
        if block.strip() and sections:
            sec = sections[-1]
            sec["code"] = (sec["code"] + "\n\n" + block) if sec["code"] else block

    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith("```"):
            if in_code:
                close_code()
                in_code, code = False, []
            else:
                in_code, code = True, []
            continue
        if in_code:
            code.append(raw.rstrip())
            continue
        if not line:
            continue
        h = re.match(r"^#{2,}\s+(.*)", line)
        t = re.match(r"^#\s+(.*)", line)
        b = re.match(r"^(?:[-*\u2022]|\d+[.)])\s+(.*)", line)
        if h:
            sections.append({"heading": _clean_line(h.group(1)) or "Section", "points": [], "code": ""})
        elif t:
            title = title or _clean_line(t.group(1))
        elif b and sections:
            point = _clean_line(b.group(1))
            if point:
                sections[-1]["points"].append(point)
    if in_code and code:
        close_code()
    sections = [sec for sec in sections if sec["points"] or sec["code"]]
    if not sections:
        return None
    return {"title": title or "StudyBuddy Notes", "sections": sections}


def to_outline(data: dict) -> str:
    out = [f"# {data['title']}"]
    for sec in data["sections"]:
        out.append(f"## {sec['heading']}")
        out += [f"- {p}" for p in sec["points"]]
        if sec.get("code"):
            out += ["```", sec["code"], "```"]
    return "\n".join(out)


def generate_document(provider, model, api_key, source_kind, source_text,
                      n_sections, level_hint, attachments=None) -> dict:
    if source_kind == "topic":
        task = ("Create study material for the request below. Follow EVERY detail of it "
                f"(topic, number of slides, examples, language, level):\n\"{source_text}\"")
    elif source_kind == "answer":
        task = f"Turn the following text into structured study material:\n\n{source_text}"
    else:
        task = "Create study material from the attached files."
    prompt = f"{task}\n\nUse exactly {n_sections} sections. {level_hint}\n{OUTLINE_FORMAT}"
    for _ in range(2):  # one automatic retry
        text = generate_text(provider, prompt, OUTLINE_SYSTEM, model, api_key, attachments)
        data = parse_outline(text or "")
        if data:
            return data
    raise ValueError("The AI did not return a usable file. Please try again.")


def revise_document(provider, model, api_key, data, instruction, level_hint) -> dict:
    """Change an existing document exactly as the student asked (add code, add slides, ...)."""
    prompt = (
        "Here is the student's current document in outline format:\n\n"
        f"{to_outline(data)}\n\n"
        f"The student now asks: \"{instruction}\"\n\n"
        "Make EXACTLY this change. Keep everything else the same unless the request says "
        f"otherwise. {level_hint}\nReturn the COMPLETE updated document.\n{OUTLINE_FORMAT}"
    )
    for _ in range(2):
        text = generate_text(provider, prompt, OUTLINE_SYSTEM, model, api_key, None)
        new = parse_outline(text or "")
        if new:
            return new
    raise ValueError("The AI did not return a usable file. Please try again.")


# ---------------- Is the student asking to CHANGE the file we just made? ----------------
_EDIT = re.compile(
    r"\b(add|include|insert|put|mention|append|remove|delete|drop|change|replace|update|modify|"
    r"edit|rewrite|shorten|shorter|longer|expand|extend|improve|redo|fix|correct|more|less|also|"
    r"translate|simplify|elaborate|detail|detailed|examples?|code|coding|programs?|programming)\b")
_EDIT_THING = re.compile(
    r"\b(slides?|ppt|pptx|presentation|pdf|document|doc|file|section|points?|bullets?|examples?|"
    r"code|coding|programs?|programming|diagram|table|summary|heading|title|it|this|them|these|"
    r"same|above)\b")
_REFERS_BACK = re.compile(r"\b(it|this|same|above|these|them)\b")
_NEW_TOPIC = re.compile(r"\b(on|about)\s+\w+")


def is_edit_request(text: str, has_last_doc: bool, last_was_file: bool) -> bool:
    low = (text or "").lower()
    if not has_last_doc:
        return False
    if detect_file_request(low) and _NEW_TOPIC.search(low) and not _REFERS_BACK.search(low):
        return False  # "make a PPT on <new topic>" is a NEW file
    mentions_file = bool(re.search(r"\b(slides?|ppt|pptx|presentation|pdf|document|file)\b", low))
    if not (last_was_file or mentions_file):
        return False
    if detect_file_request(low) and _REFERS_BACK.search(low):
        return True  # "make it a PDF" / "convert this to Word"
    return bool(_EDIT.search(low) and _EDIT_THING.search(low))


# ---------------- Build the files (PowerPoint / Word / PDF) with code boxes ----------------
def _code_lines(code: str) -> list:
    lines = (code or "").expandtabs(4).split("\n")
    while lines and not lines[-1].strip():
        lines.pop()
    return lines


def build_pptx(data: dict) -> bytes:
    import io
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    NAVY = RGBColor(0x1F, 0x2A, 0x44)
    ACCENT = RGBColor(0x4F, 0x8E, 0xF7)
    LIGHT = RGBColor(0xF5, 0xF7, 0xFB)
    TEXT = RGBColor(0x22, 0x2B, 0x3A)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    SOFT = RGBColor(0xC8, 0xD3, 0xEA)
    CODE_BG = RGBColor(0x1E, 0x1E, 0x2E)
    CODE_FG = RGBColor(0xE6, 0xEA, 0xF5)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def rect(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    def textbox(slide, x, y, w, h, text, size, color, bold=False, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        run = frame.paragraphs[0].add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
        return box

    def chunks(seq, n):
        return [seq[i:i + n] for i in range(0, len(seq), n)]

    # Title slide
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, NAVY)
    textbox(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.0), data["title"],
            44, WHITE, True, MSO_ANCHOR.BOTTOM)
    rect(s, Inches(0.8), Inches(3.8), Inches(1.6), Inches(0.08), ACCENT)
    textbox(s, Inches(0.8), Inches(4.1), Inches(11), Inches(0.8), "Created with StudyBuddy", 20, SOFT)

    for sec in data["sections"]:
        code = _code_lines(sec.get("code", ""))
        bullet_chunks = chunks(sec["points"], 4 if code else 6) or [[]]
        code_chunks = chunks(code, 14) if code else [[]]
        for idx in range(max(len(bullet_chunks), len(code_chunks))):
            bullets = bullet_chunks[idx] if idx < len(bullet_chunks) else []
            snippet = code_chunks[idx] if idx < len(code_chunks) else []
            s = prs.slides.add_slide(blank)
            rect(s, 0, 0, W, H, LIGHT)
            rect(s, 0, 0, W, Inches(1.3), NAVY)
            rect(s, 0, Inches(1.3), W, Inches(0.07), ACCENT)
            heading = sec["heading"] + (" (cont.)" if idx else "")
            textbox(s, Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.9),
                    heading, 30, WHITE, True, MSO_ANCHOR.MIDDLE)

            two_cols = bool(bullets and snippet)
            if bullets:
                longest = max(len(p) for p in bullets)
                if two_cols:
                    size = 20 if longest <= 70 else 17 if longest <= 110 else 15
                    width = Inches(5.6)
                else:
                    size = 22 if longest <= 90 else 19 if longest <= 150 else 16
                    width = Inches(11.5)
                body = s.shapes.add_textbox(Inches(0.9), Inches(1.8), width, Inches(5.2))
                frame = body.text_frame
                frame.word_wrap = True
                for n, point in enumerate(bullets):
                    para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                    para.space_after = Pt(12)
                    run = para.add_run()
                    run.text = "\u2022  " + point
                    run.font.size, run.font.color.rgb = Pt(size), TEXT
            if snippet:
                x = Inches(6.8) if two_cols else Inches(0.9)
                w = Inches(5.9) if two_cols else Inches(11.5)
                rect(s, x, Inches(1.8), w, Inches(5.1), CODE_BG)
                widest = max(len(line) for line in snippet)
                if two_cols:
                    size = 16 if widest <= 36 else 14 if widest <= 44 else 12 if widest <= 54 else 11
                else:
                    size = 20 if widest <= 60 else 16 if widest <= 80 else 13
                box = s.shapes.add_textbox(x + Inches(0.2), Inches(1.95), w - Inches(0.4), Inches(4.8))
                frame = box.text_frame
                frame.word_wrap = True
                for n, line in enumerate(snippet):
                    para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                    run = para.add_run()
                    run.text = line if line else " "
                    run.font.name = "Consolas"
                    run.font.size, run.font.color.rgb = Pt(size), CODE_FG

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_docx(data: dict) -> bytes:
    import io
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(data["title"], 0)
    for sec in data["sections"]:
        doc.add_heading(sec["heading"], level=1)
        for point in sec["points"]:
            doc.add_paragraph(point, style="List Bullet")
        if sec.get("code"):
            para = doc.add_paragraph()
            run = para.add_run("\n".join(_code_lines(sec["code"])))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:color"), "auto")
            shade.set(qn("w:fill"), "F0F2F5")
            para._p.get_or_add_pPr().append(shade)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _latin(text: str) -> str:
    """The basic PDF fonts only support Latin letters, so clean the text."""
    for a, b in {"\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"',
                 "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-"}.items():
        text = text.replace(a, b)
    text = text.encode("latin-1", "replace").decode("latin-1")
    return re.sub(r"(\S{60})(?=\S)", r"\1 ", text)  # break very long words


def build_pdf(data: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 20)
    pdf.multi_cell(0, 10, _latin(data["title"]), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for sec in data["sections"]:
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(0, 8, _latin(sec["heading"]), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for point in sec["points"]:
            pdf.multi_cell(0, 6, _latin("- " + point), new_x="LMARGIN", new_y="NEXT")
        if sec.get("code"):
            pdf.ln(1)
            pdf.set_font("Courier", "", 9)
            pdf.set_fill_color(240, 242, 245)
            for line in _code_lines(sec["code"]):
                pdf.multi_cell(0, 5, _latin(line) or " ", new_x="LMARGIN", new_y="NEXT", fill=True)
        pdf.ln(3)
    return bytes(pdf.output())


FORMATS = {
    "PowerPoint (.pptx)": (build_pptx, "pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "PDF (.pdf)": (build_pdf, "pdf", "application/pdf"),
    "Word (.docx)": (build_docx, "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
}


# ---------------- Current information: today's date + automatic web search ----------------
ANSWER_RULES = (
    "Quality rules: give accurate, well-organised answers in simple language. "
    "For history, give dates, causes and effects. "
    "For current affairs or news, explain the background (past), the present situation and what "
    "may happen next (future), and show the different viewpoints fairly. "
    "Your own knowledge can be out of date: for anything recent, rely on the web results given "
    "below and mention the source name and date. If web results are missing or unclear, say "
    "clearly that you could not confirm the latest information instead of guessing. "
    "Never invent facts, names, numbers, dates or links. If you are unsure, say so. "
    "Follow EVERY instruction in the student's message exactly (number of items, language, "
    "examples, code, format). If a follow-up message changes or adds to your previous answer, "
    "update that answer instead of starting something new."
)

_WEB_WORDS = re.compile(
    r"\b(news|latest|current|currently|today|tonight|yesterday|tomorrow|recent|recently|"
    r"breaking|trending|this (week|month|year)|last (week|month|year)|"
    r"election|elections|budget|stock price|sensex|nifty|score|ipl|world cup|weather|"
    r"ceasefire|upcoming|what happened|what's happening|going on|who won|20(2[4-9]|3\d))\b",
    re.I)


def needs_web(text: str) -> bool:
    """True when the question is about recent / changing information."""
    return bool(_WEB_WORDS.search(text or ""))


def today_text() -> str:
    from datetime import datetime, timedelta, timezone
    ist = datetime.now(timezone(timedelta(hours=5, minutes=30)))
    return ist.strftime("%A, %d %B %Y")


def search_current(query: str, n: int = 6) -> list:
    """News + web search (DuckDuckGo). Returns [] if anything goes wrong."""
    found = []
    try:
        try:
            from ddgs import DDGS
        except ImportError:  # older package name
            from duckduckgo_search import DDGS
        d = DDGS()
        try:
            for r in d.news(query, max_results=4):
                day = (r.get("date") or "")[:10]
                found.append({"title": r.get("title", ""), "url": r.get("url", ""),
                              "snippet": ((day + " - ") if day else "") + (r.get("body") or "")})
        except Exception:
            pass
        try:
            for r in d.text(query, max_results=n):
                found.append({"title": r.get("title", ""), "url": r.get("href", ""),
                              "snippet": r.get("body") or ""})
        except Exception:
            pass
    except Exception:
        pass
    if not found:
        found = search_web(query)
    seen, out = set(), []
    for r in found:
        if r["url"] and r["url"] not in seen:
            seen.add(r["url"])
            r["snippet"] = r["snippet"][:300]
            out.append(r)
    return out[: n + 2]


# ---------------- 3. Understand "make me a PPT on ..." typed in the chat ----------------
_MAKE = re.compile(r"\b(make|create|generate|build|prepare|produce|convert|turn|export|"
                   r"give|write|draft|need|want|download|get)\b")
_START = re.compile(r"^\s*(a |an )?(pptx?|power\s?point|pdf|presentation|slides?|docx?)\s+"
                    r"(on|about|for|of)\b")
_FORMAT_WORDS = {
    "PowerPoint (.pptx)": re.compile(r"\b(pptx?|power\s?point|slides?|slide\s?deck|presentation|deck)\b"),
    "PDF (.pdf)": re.compile(r"\bpdf\b"),
    "Word (.docx)": re.compile(r"\b(docx|ms\s?word|word\s+(file|document|doc)|"
                               r"(in|as|to|into)\s+(a\s+)?word)\b"),
}
# "...from THIS pdf" means the pdf is the source, not the file we should make
_SOURCE_BEFORE = re.compile(r"\b(this|the|attached|uploaded|my|from|of|read|summari[sz]e)\s+$")
_ANSWER_WORDS = re.compile(r"\b(above|previous|last answer|this answer|that answer|your answer|"
                           r"your last|the answer|these notes|this explanation)\b")
_FILE_WORDS = re.compile(r"\b(attach\w*|upload\w*|this (file|pdf|doc\w*|image|photo|notes)|"
                         r"these (files|notes)|the (file|pdf|document)|my (file|notes|pdf))\b")


def detect_file_request(text: str):
    """Return 'PowerPoint (.pptx)', 'PDF (.pdf)' or 'Word (.docx)' if the student
    is asking us to MAKE such a file, otherwise None."""
    low = text.lower()
    if not (_MAKE.search(low) or _START.search(low)):
        return None
    found = []
    for fmt, pat in _FORMAT_WORDS.items():
        for m in pat.finditer(low):
            before = low[max(0, m.start() - 14):m.start()]
            if _SOURCE_BEFORE.search(before) and not re.match(r"(in|as|to|into)\s", m.group(0)):
                continue
            found.append((m.start(), fmt))
    return min(found)[1] if found else None


def plan_file_request(text: str, has_files: bool, has_answer: bool) -> dict:
    """Decide what the file should be made from and how many slides/sections."""
    low = text.lower()
    n = re.search(r"\b(\d{1,2})\s*(?:slides?|pages?|sections?)\b", low)
    n_sections = min(12, max(3, int(n.group(1)))) if n else 6
    if has_answer and _ANSWER_WORDS.search(low):
        return {"kind": "answer", "n": n_sections}
    if _FILE_WORDS.search(low):
        return {"kind": "files" if has_files else "need_file", "n": n_sections}
    return {"kind": "topic", "n": n_sections}


# ---------------- Default settings (change them here, never shown on screen) ----------------
PROVIDER = "gemini"           # "gemini" (online) or "ollama" (offline)
MODEL = "gemini-3.6-flash"    # for ollama use e.g. "llama3.1:8b"
TONE = "Friendly Tutor"       # Friendly Tutor / Strict Examiner / Fun Mode
LEVEL = "School"              # Like I'm 10 / School / College / Expert
DEEP_THINK = False
MAX_HISTORY = 20              # only send the last N messages to the model

st.set_page_config(
    page_title="StudyBuddy - Student AI Assistant",
    page_icon="🎓",
    layout="centered",
    initial_sidebar_state="collapsed",
    menu_items={"Get help": None, "Report a bug": None, "About": None},
)

# ---------------- Look and feel: dark space theme ----------------
CSS = """
<style>
:root {
  --sb-text: #e8ecff; --sb-muted: #9aa6d6; --sb-accent: #6c8cff; --sb-accent2: #8b5cf6;
  --sb-border: rgba(130,150,255,.24); --sb-glass: rgba(12,18,52,.62);
}
html, body { color-scheme: dark; }

/* Deep black-blue galaxy with subtle stars */
.stApp {
  color: var(--sb-text);
  background-color: #02030a;
  background-image:
    radial-gradient(circle at 20% 30%, rgba(255,255,255,.90) 0, rgba(255,255,255,0) 1.4px),
    radial-gradient(circle at 70% 80%, rgba(200,215,255,.75) 0, rgba(255,255,255,0) 1.2px),
    radial-gradient(circle at 90% 15%, rgba(255,255,255,.90) 0, rgba(255,255,255,0) 1.6px),
    radial-gradient(circle at 10% 70%, rgba(255,255,255,.70) 0, rgba(255,255,255,0) 1.2px),
    radial-gradient(circle at 55% 25%, rgba(190,205,255,.80) 0, rgba(255,255,255,0) 1.4px),
    radial-gradient(circle at 85% 60%, rgba(255,255,255,.65) 0, rgba(255,255,255,0) 1.2px),
    radial-gradient(circle at 35% 55%, rgba(255,255,255,.60) 0, rgba(255,255,255,0) 1.1px),
    radial-gradient(circle at 65% 10%, rgba(255,255,255,.85) 0, rgba(255,255,255,0) 1.5px),
    radial-gradient(circle at 5% 95%, rgba(210,225,255,.70) 0, rgba(255,255,255,0) 1.2px),
    radial-gradient(ellipse 80% 55% at 12% -5%, rgba(58,84,210,.38), transparent 70%),
    radial-gradient(ellipse 70% 50% at 100% 105%, rgba(96,52,190,.28), transparent 70%),
    linear-gradient(180deg, #01020a 0%, #040a22 55%, #071340 100%);
  background-size: 260px 260px, 260px 260px, 260px 260px,
                   340px 340px, 340px 340px, 340px 340px,
                   470px 470px, 470px 470px, 470px 470px,
                   100% 100%, 100% 100%, 100% 100%;
  background-attachment: fixed;
}
[data-testid="stAppViewContainer"], [data-testid="stMain"], .main,
[data-testid="stHeader"] { background: transparent !important; }

/* Hide Streamlit menu, deploy button and footer (the sidebar button stays visible) */
#MainMenu, footer, [data-testid="stToolbarActions"], [data-testid="stMainMenu"],
[data-testid="stAppDeployButton"], .stAppDeployButton, .stDeployButton,
[data-testid="stDecoration"], [data-testid="stStatusWidget"] { display: none !important; }

/* History button (top-left): show a 3-line menu icon */
[data-testid="stSidebarCollapsedControl"], [data-testid="collapsedControl"],
[data-testid="stExpandSidebarButton"] { visibility: visible !important; z-index: 1000002; }
[data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapsedControl"] button,
[data-testid="collapsedControl"] button { position: relative; }
[data-testid="stExpandSidebarButton"] svg, [data-testid="stExpandSidebarButton"] [data-testid="stIconMaterial"],
[data-testid="stSidebarCollapsedControl"] button svg, [data-testid="stSidebarCollapsedControl"] [data-testid="stIconMaterial"],
[data-testid="collapsedControl"] button svg, [data-testid="collapsedControl"] [data-testid="stIconMaterial"] {
  opacity: 0 !important; }
[data-testid="stExpandSidebarButton"]::after, [data-testid="stSidebarCollapsedControl"] button::after,
[data-testid="collapsedControl"] button::after {
  content: "☰"; position: absolute; inset: 0; display: flex; align-items: center;
  justify-content: center; font-size: 1.5rem; color: #e8ecff; pointer-events: none; }

/* History sidebar */
[data-testid="stSidebar"] { background: rgba(5,8,30,.97) !important;
  border-right: 1px solid var(--sb-border); }
[data-testid="stSidebar"] .stButton button { justify-content: flex-start; text-align: left; }
[data-testid="stSidebar"] .stButton button p { white-space: nowrap; overflow: hidden;
  text-overflow: ellipsis; }
.stButton button[data-testid="stBaseButton-primary"] { background: rgba(108,140,255,.28);
  border-color: var(--sb-accent); }
.sb-side-title { font-weight: 700; font-size: 1.05rem; margin: 0 0 .4rem; }

.block-container { max-width: 860px; padding: 2rem 1.2rem 8rem; }

/* Header */
.sb-hero { text-align: center; margin: .3rem 0 1.3rem; }
.sb-logo { font-size: 2.4rem; filter: drop-shadow(0 0 14px rgba(108,140,255,.75)); }
.sb-title { font-size: 2.5rem; font-weight: 800; letter-spacing: -.02em; margin: .1rem 0 .2rem;
  background: linear-gradient(90deg,#ffffff,#9db4ff 55%,#c4a8ff);
  -webkit-background-clip: text; background-clip: text; -webkit-text-fill-color: transparent; }
.sb-sub { color: var(--sb-muted); font-size: 1rem; max-width: 560px; margin: 0 auto; }

/* Buttons */
.stButton button, .stDownloadButton button {
  background: var(--sb-glass); color: var(--sb-text); border: 1px solid var(--sb-border);
  border-radius: 12px; padding: .5rem .9rem; backdrop-filter: blur(8px);
  transition: all .18s ease; }
.stButton button:hover, .stDownloadButton button:hover {
  color: #fff; border-color: var(--sb-accent); transform: translateY(-1px);
  box-shadow: 0 6px 20px rgba(60,90,230,.28); }

/* Panels and inputs */
[data-testid="stExpander"] { background: var(--sb-glass); border: 1px solid var(--sb-border);
  border-radius: 14px; backdrop-filter: blur(8px); }
[data-testid="stExpander"] details { border: none !important; background: transparent !important; }
.stTextInput input, [data-baseweb="select"] > div, [data-baseweb="input"] {
  background: rgba(8,12,40,.72) !important; color: var(--sb-text) !important;
  border-color: var(--sb-border) !important; border-radius: 10px !important; }
[data-testid="stAlert"] { border-radius: 12px; backdrop-filter: blur(6px); }
pre, [data-testid="stCode"] pre { background: rgba(4,8,30,.88) !important;
  border: 1px solid var(--sb-border); border-radius: 10px; }
.stMarkdown table { display: block; overflow-x: auto; }
img { max-width: 100%; }

/* Chat messages: clearly different for You and the AI */
.sb-marker { display: none; }
[data-testid="stElementContainer"]:has(> [data-testid="stMarkdown"] .sb-marker),
.element-container:has(> .stMarkdown .sb-marker) { display: none !important; }
[data-testid="stChatMessage"] { border-radius: 18px; padding: .85rem 1.05rem; margin: .55rem 0;
  border: 1px solid rgba(255,255,255,.08); backdrop-filter: blur(8px); }
[data-testid="stChatMessage"]:has(.sb-user) {
  background: linear-gradient(135deg, rgba(79,110,255,.34), rgba(124,77,255,.24));
  border-color: rgba(130,150,255,.42); flex-direction: row-reverse; margin-left: 12%;
  border-bottom-right-radius: 6px; }
[data-testid="stChatMessage"]:has(.sb-ai) {
  background: rgba(255,255,255,.045); border-color: rgba(255,255,255,.11);
  margin-right: 6%; border-bottom-left-radius: 6px;
  box-shadow: inset 3px 0 0 rgba(108,140,255,.55); }

/* Thinking indicator */
.sb-thinking { display: inline-flex; align-items: center; gap: .45rem; color: var(--sb-muted);
  font-style: italic; padding: .2rem 0; }
.sb-thinking .dot { width: 7px; height: 7px; border-radius: 50%; background: #8fa8ff;
  animation: sb-bounce 1.2s infinite ease-in-out both; }
.sb-thinking .dot:nth-child(2) { animation-delay: .15s; }
.sb-thinking .dot:nth-child(3) { animation-delay: .30s; }
@keyframes sb-bounce { 0%,80%,100% { transform: scale(.5); opacity: .4; }
                       40% { transform: scale(1); opacity: 1; } }

/* Input box and send button */
[data-testid="stBottom"], [data-testid="stBottom"] > div, [data-testid="stBottomBlockContainer"] {
  background: transparent !important; }
[data-testid="stBottom"] { background: linear-gradient(180deg, transparent, rgba(2,3,10,.88) 45%) !important; }
[data-testid="stChatInput"] { background: rgba(10,16,48,.88); border: 1px solid rgba(120,140,255,.40);
  border-radius: 28px; box-shadow: 0 8px 30px rgba(0,0,0,.5), 0 0 24px rgba(70,100,255,.14);
  transition: border-color .2s, box-shadow .2s; }
[data-testid="stChatInput"]:focus-within { border-color: #7d9bff;
  box-shadow: 0 0 0 3px rgba(108,140,255,.24), 0 8px 30px rgba(0,0,0,.5); }
[data-testid="stChatInput"] > div { background: transparent !important; border: none !important; }
[data-testid="stChatInput"] textarea { background: transparent !important; color: var(--sb-text) !important;
  caret-color: #9db4ff; font-size: 1rem; }
[data-testid="stChatInput"] textarea::placeholder { color: rgba(170,185,235,.62); }
[data-testid="stChatInput"] button { border-radius: 50% !important; }
[data-testid="stChatInputSubmitButton"] { background: linear-gradient(135deg,#5b7cff,#8b5cf6) !important;
  color: #fff !important; box-shadow: 0 4px 14px rgba(91,124,255,.45); }
[data-testid="stChatInputSubmitButton"]:disabled { opacity: .45; box-shadow: none; }

/* Mobile */
@media (max-width: 640px) {
  .block-container { padding: 1.2rem .7rem 8rem; }
  .sb-title { font-size: 1.85rem; } .sb-sub { font-size: .92rem; }
  [data-testid="stChatMessage"]:has(.sb-user) { margin-left: 4%; }
  [data-testid="stChatMessage"]:has(.sb-ai) { margin-right: 0; }
  [data-testid="stHorizontalBlock"] { flex-wrap: wrap !important; gap: .5rem !important; }
  [data-testid="stColumn"], [data-testid="column"] {
    min-width: calc(50% - .5rem) !important; flex: 1 1 calc(50% - .5rem) !important; }
}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

USER_MARK = '<span class="sb-marker sb-user"></span>'
AI_MARK = '<span class="sb-marker sb-ai"></span>'


def mark(role: str):
    st.markdown(USER_MARK if role == "user" else AI_MARK, unsafe_allow_html=True)


def thinking_html(label: str) -> str:
    return ('<div class="sb-thinking"><span class="dot"></span><span class="dot"></span>'
            f'<span class="dot"></span><span>{label}…</span></div>')


# ---------------- Errors: details go to the CMD window, never to the page ----------------
SAFE_PREFIXES = ("The AI did not", "The AI returned", "The assistant is not", "Google could not")


def friendly_error(e: Exception) -> str:
    print("\n[StudyBuddy] An error occurred (shown here only, never on the web page):",
          file=sys.stderr)
    traceback.print_exception(type(e), e, e.__traceback__)
    msg = str(e)
    if isinstance(e, ValueError) and msg.startswith(SAFE_PREFIXES):
        return msg
    low = msg.lower()
    if any(k in low for k in ("429", "quota", "resource_exhausted", "rate limit")):
        return "Too many requests right now. Please wait a minute and try again."
    if any(k in low for k in ("api key", "api_key", "401", "403", "permission", "unauthenticated")):
        return "The assistant is not available right now. Please try again later."
    if any(k in low for k in ("404", "not_found", "no longer available")):
        return "The assistant is temporarily unavailable. Please try again later."
    if any(k in low for k in ("timeout", "timed out", "connect", "getaddrinfo", "network",
                              "503", "unavailable", "overloaded")):
        return "I couldn't reach the AI service. Please check your internet and try again."
    return "Something went wrong. Please try again."


def get_key() -> str:
    try:
        return st.secrets["GEMINI_API_KEY"]
    except Exception:
        return os.environ.get("GEMINI_API_KEY", "")


api_key = get_key()

# ---------------- Session state ----------------
ss = st.session_state
if not ss.get("db_ready"):  # set up the history database once, not on every click
    db.init()
    ss.db_ready = True
ss.setdefault("messages", [])      # chat shown on screen
ss.setdefault("attachments", [])   # files attached to this chat
ss.setdefault("chat_id", None)     # id in the saved-history database
ss.setdefault("queued", None)      # prompt from a quick-action button
ss.setdefault("last_audio", None)  # to avoid answering the same voice note twice
ss.setdefault("generated", None)   # last PPT / PDF / Word file made


def last_answer() -> str:
    for m in reversed(ss.messages):
        if m["role"] == "assistant" and not m.get("file") and not m["content"].startswith("📎"):
            return m["content"]
    return ""


def make_file_message(request: str, fmt: str, edit: bool = False) -> dict:
    """Build (or change) a PPT / PDF / Word file from what the student typed in the chat."""
    last = ss.get("last_doc")
    verb = "ready"
    if edit and last:
        if _EDIT.search(request.lower()):
            data = revise_document(PROVIDER, MODEL, api_key, last["data"], request, LEVELS[LEVEL])
            verb = "updated"
        else:  # only a format change, e.g. "make it a PDF"
            data = last["data"]
            verb = "converted"
    else:
        plan = plan_file_request(request, bool(ss.attachments), bool(last_answer()))
        if plan["kind"] == "need_file":
            return {"role": "assistant",
                    "content": "Please attach your file first with the **+** button in the chat box, "
                               "then ask me again."}
        kind = plan["kind"]
        source_text = last_answer() if kind == "answer" else request
        data = generate_document(PROVIDER, MODEL, api_key, kind, source_text, plan["n"],
                                 LEVELS[LEVEL], ss.attachments if kind == "files" else None)
    builder, ext, mime = FORMATS[fmt]
    name = f"{safe_name(data['title'])}.{ext}"
    ss.last_doc = {"data": data, "fmt": fmt}
    outline = "\n".join(f"{i}. {sec['heading']}" + (" 💻" if sec.get("code") else "")
                        for i, sec in enumerate(data["sections"], 1))
    label = fmt.split(" (")[0]
    return {"role": "assistant",
            "content": f"📎 Your {label} file is {verb}: **{name}**\n\n**Contents:**\n{outline}",
            "file": {"name": name, "bytes": builder(data), "mime": mime}}


def reset_chat():
    ss.messages, ss.attachments, ss.chat_id, ss.generated = [], [], None, None
    ss.last_doc = None


def open_chat(cid):
    reset_chat()
    ss.chat_id = cid
    ss.messages = [{"role": r, "content": c} for r, c in db.load_messages(cid)]


def delete_current_chat():
    if ss.chat_id is not None:
        db.delete_chat(ss.chat_id)
    reset_chat()


def render_message(m, idx=0):
    with st.chat_message(m["role"]):
        mark(m["role"])
        for f in m.get("files", []):
            if f["kind"] == "image" and f.get("data"):
                st.image(f["data"], width=260)
            else:
                st.caption(f"📎 {f['name']}")
        st.markdown(m["content"])
        made = m.get("file")
        if made:
            st.download_button("⬇️ Download " + made["name"], made["bytes"],
                               file_name=made["name"], mime=made["mime"], key=f"dl_{idx}")


# ---------------- Sidebar: chat history (open it with the 3-line button, top-left) ----------------
with st.sidebar:
    st.markdown('<div class="sb-side-title">🎓 StudyBuddy</div>', unsafe_allow_html=True)
    st.button("➕ New chat", use_container_width=True, on_click=reset_chat, key="new_chat")
    st.caption("Chat history")
    saved_chats = db.list_chats(50)
    if not saved_chats:
        st.caption("No saved chats yet. Your chats will appear here.")
    for cid, title in saved_chats:
        st.button("💬 " + title, key=f"chat_{cid}", use_container_width=True,
                  type="primary" if cid == ss.chat_id else "secondary",
                  on_click=open_chat, args=(cid,))
    if ss.chat_id is not None:
        st.button("🗑️ Delete this chat", use_container_width=True,
                  on_click=delete_current_chat, key="del_chat")


# ---------------- Header ----------------
st.markdown(
    '<div class="sb-hero"><div class="sb-logo">🎓</div><div class="sb-title">StudyBuddy</div>'
    '<div class="sb-sub">Your free AI study assistant. Ask anything, attach files, '
    'and turn answers into slides, PDFs and notes.</div></div>',
    unsafe_allow_html=True,
)

if PROVIDER == "gemini" and not api_key:
    if not ss.get("warned_key"):
        print("[StudyBuddy] GEMINI_API_KEY is missing. Add it to .streamlit/secrets.toml",
              file=sys.stderr)
        ss.warned_key = True
    st.error("The assistant is not available right now. Please try again later.")
    st.stop()

# Quick study tools
cols = st.columns(len(QUICK_ACTIONS))
for col, (label, text) in zip(cols, QUICK_ACTIONS.items()):
    if col.button(label, use_container_width=True):
        ss.queued = text

use_web = st.toggle("🌐 Always search the web (it also searches by itself for news and current topics)")

# ---------------- Tools panels ----------------
with st.expander("📄 Create a PPT / PDF / Word file"):
    source = st.radio("Make it from",
                      ["A topic I type", "My last chat answer", "My attached files"],
                      horizontal=True)
    topic = st.text_input("Topic", placeholder="e.g. Photosynthesis") \
        if source == "A topic I type" else ""
    c1, c2 = st.columns(2)
    fmt = c1.selectbox("Format", list(FORMATS))
    n_sections = c2.slider("Slides / sections", 3, 12, 6)
    if fmt.startswith("PDF"):
        st.caption("PDF supports English letters only. Use PowerPoint or Word for other languages.")

    if st.button("✨ Create file"):
        problem = None
        if source == "A topic I type" and not topic.strip():
            problem = "Please type a topic first."
        elif source == "My last chat answer" and not last_answer():
            problem = "There is no chat answer yet. Ask something first."
        elif source == "My attached files" and not ss.attachments:
            problem = "No files attached yet. Use the + button in the chat box."
        if problem:
            st.warning(problem)
        else:
            kind = {"A topic I type": "topic", "My last chat answer": "answer"}.get(source, "files")
            text = topic.strip() if kind == "topic" else last_answer()
            with st.spinner("Creating your file... this can take up to a minute"):
                try:
                    data = generate_document(
                        PROVIDER, MODEL, api_key, kind, text, n_sections,
                        LEVELS[LEVEL], ss.attachments if kind == "files" else None)
                    builder, ext, mime = FORMATS[fmt]
                    ss.generated = {"name": f"{safe_name(data['title'])}.{ext}",
                                    "bytes": builder(data), "mime": mime}
                    ss.last_doc = {"data": data, "fmt": fmt}
                except Exception as e:
                    ss.generated = None
                    st.error(friendly_error(e))

    if ss.generated:
        st.success("Your file is ready.")
        st.download_button("⬇️ Download " + ss.generated["name"], ss.generated["bytes"],
                           file_name=ss.generated["name"], mime=ss.generated["mime"])

with st.expander("🎤 Ask by voice"):
    audio = st.audio_input("Record your question") if hasattr(st, "audio_input") else None
    if not hasattr(st, "audio_input") or PROVIDER != "gemini":
        st.caption("Voice input is not available right now.")

if ss.attachments:
    st.caption("📎 Attached to this chat: " + ", ".join(a["name"] for a in ss.attachments))
    if PROVIDER == "ollama" and any(not a["text"] for a in ss.attachments):
        st.warning("Offline mode can read text, PDF, Word and PPT files only. "
                   "Images, audio and video need the online engine.")
    if st.button("Remove attachments"):
        ss.attachments = []
        st.rerun()

# ---------------- New input ----------------
value = st.chat_input("Ask StudyBuddy anything...", accept_file="multiple", file_type=ALLOWED_TYPES)
typed, new_files = "", []
if isinstance(value, str):
    typed = value.strip()
elif value:
    typed = (value["text"] or "").strip()
    new_files = list(value["files"] or [])

# Voice note (only the first time we see a new recording)
voice_att = []
if audio is not None:
    data = audio.getvalue()
    digest = hashlib.md5(data).hexdigest()
    if digest != ss.last_audio:
        ss.last_audio = digest
        if PROVIDER == "gemini":
            voice_att = [make_attachment("voice.wav", data)]

queued, ss.queued = ss.queued, None

prompt = typed or None
shown_files = []
for f in new_files:
    att = make_attachment(f.name, f.getvalue(), read_pdf_text=(PROVIDER == "ollama"))
    if not any(a["name"] == att["name"] and len(a["data"]) == len(att["data"])
               for a in ss.attachments):
        ss.attachments.append(att)
    shown_files.append({"name": att["name"], "kind": att["kind"],
                        "data": att["data"] if att["kind"] == "image" else None})
if not prompt and new_files:
    prompt = "Please explain the attached file(s) and tell me the key points."
if not prompt and voice_att:
    prompt = ("🎤 Voice question: first write what I said in one short line, "
              "then answer it.")
if not prompt and queued:
    prompt = queued

file_fmt = detect_file_request(typed) if typed else None  # "make a PPT on ..." typed in chat
_last_msg = next((m for m in reversed(ss.messages) if m["role"] == "assistant"), None)
edit_mode = bool(typed) and is_edit_request(typed, bool(ss.get("last_doc")),
                                            bool(_last_msg and _last_msg.get("file")))
if edit_mode:  # "add programming examples" / "make it a PDF" -> change the file we already made
    file_fmt = file_fmt or ss.last_doc["fmt"]

if prompt:
    ss.messages.append({"role": "user", "content": prompt, "files": shown_files})

# ---------------- Chat: ONE container, every message drawn the same way ----------------
answered = False
chat_area = st.container()
with chat_area:
    for i, m in enumerate(ss.messages):
        render_message(m, i)

    if prompt and file_fmt:
        with st.chat_message("assistant"):
            mark("assistant")
            holder = st.empty()
            holder.markdown(thinking_html("Creating your " + file_fmt.split(" (")[0] + " file"),
                            unsafe_allow_html=True)
            try:
                made_msg = make_file_message(prompt, file_fmt, edit_mode)
                holder.empty()
                ss.messages.append(made_msg)
                try:  # saving history must never break the chat
                    if ss.chat_id is None:
                        ss.chat_id = db.new_chat(prompt)
                    db.add_message(ss.chat_id, "user", prompt)
                    db.add_message(ss.chat_id, "assistant", made_msg["content"])
                except Exception:
                    traceback.print_exc()
                answered = True
            except Exception as e:
                holder.empty()
                st.error(friendly_error(e))
                ss.messages.pop()  # remove the unanswered question

    elif prompt:
        with st.chat_message("assistant"):
            mark("assistant")
            holder = st.empty()  # the "Thinking..." indicator lives here until the first words arrive
            busy = "Reading your files" if (ss.attachments or voice_att) else "Thinking"
            holder.markdown(thinking_html(busy), unsafe_allow_html=True)
            try:
                system = build_system_prompt(TONE, LEVEL, DEEP_THINK)
                system += f"\n\nToday's date is {today_text()}.\n{ANSWER_RULES}"
                sources = []
                if typed and (use_web or needs_web(typed)):
                    holder.markdown(thinking_html("Searching the web"), unsafe_allow_html=True)
                    sources = search_current(typed)
                    if sources:
                        system += ("\n\nWeb search results (reference only, may be incomplete or wrong; "
                                   "mention the source name and date when you use them):\n" + "\n".join(
                                       f"- {r['title']}: {r['snippet']} ({r['url']})" for r in sources))
                    else:
                        system += ("\n\nA web search was tried but returned nothing. If the question "
                                   "needs the latest information, say you could not confirm it.")
                    holder.markdown(thinking_html(busy), unsafe_allow_html=True)

                history = [{"role": m["role"], "content": m["content"]}
                           for m in ss.messages[-MAX_HISTORY:]]
                stream = stream_reply(PROVIDER, history, system, MODEL, api_key,
                                      ss.attachments + voice_att)
                first = next(stream, "")  # waits for the first words (file uploads happen here)
                holder.empty()

                def pieces():
                    yield first
                    yield from stream

                reply = st.write_stream(pieces())
                if not reply.strip():
                    raise ValueError("The AI did not return an answer. Please try rephrasing your question.")
                if sources:
                    src = "\n\n**Sources:**\n" + "\n".join(
                        f"- [{r['title']}]({r['url']})" for r in sources if r["url"])
                    st.markdown(src)
                    reply += src
                ss.messages.append({"role": "assistant", "content": reply})

                try:  # saving history must never break the chat
                    if ss.chat_id is None:
                        ss.chat_id = db.new_chat(prompt)
                    db.add_message(ss.chat_id, "user", prompt)
                    db.add_message(ss.chat_id, "assistant", reply)
                except Exception:
                    traceback.print_exc()
                answered = True
            except Exception as e:
                holder.empty()
                st.error(friendly_error(e))
                ss.messages.pop()  # remove the unanswered question

# After a good answer, redraw once from saved messages so old answers stay put
# and the new chat shows up in the sidebar history.
if answered:
    st.rerun()
