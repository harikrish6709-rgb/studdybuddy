"""Makes PowerPoint, PDF and Word files. The AI writes the content as JSON,
then Python libraries build the actual file."""
import io
import json
import re

from llm import generate_text

DOC_SYSTEM = (
    "You write clear, accurate study material for students. "
    "Reply with ONLY valid JSON. No explanations, no markdown fences."
)


# ---------------- 1. Ask the AI for the content ----------------
def _parse_json(text: str) -> dict:
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("The AI did not return a valid file layout. Please try again.")
    try:
        return json.loads(text[start:end + 1])
    except json.JSONDecodeError:
        raise ValueError("The AI did not return a valid file layout. Please try again.")


def _normalize(data: dict) -> dict:
    title = str(data.get("title") or "StudyBuddy Notes")
    sections = []
    for sec in data.get("sections", []):
        heading = str(sec.get("heading") or "Section")
        points = [str(p) for p in sec.get("points", []) if str(p).strip()]
        sections.append({"heading": heading, "points": points})
    if not sections:
        raise ValueError("The AI returned an empty file. Please try again.")
    return {"title": title, "sections": sections}


def generate_document(provider, model, api_key, source_kind, source_text,
                      n_sections, level_hint, attachments=None) -> dict:
    if source_kind == "topic":
        task = f"Create study material on this topic: {source_text}"
    elif source_kind == "answer":
        task = f"Turn the following text into structured study material:\n\n{source_text}"
    else:
        task = "Create study material from the attached files."
    prompt = (
        f"{task}\n\nUse exactly {n_sections} sections. Each section has a short heading "
        f"and 3 to 6 short, clear points (one sentence each). {level_hint}\n"
        'Reply with ONLY JSON in this shape: {"title": "...", "sections": '
        '[{"heading": "...", "points": ["...", "..."]}]}'
    )
    text = generate_text(provider, prompt, DOC_SYSTEM, model, api_key, attachments)
    return _normalize(_parse_json(text))


# ---------------- 2. Build the files ----------------
def build_pptx(data: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    NAVY = RGBColor(0x1F, 0x2A, 0x44)
    ACCENT = RGBColor(0x4F, 0x8E, 0xF7)
    LIGHT = RGBColor(0xF5, 0xF7, 0xFB)
    TEXT = RGBColor(0x22, 0x2B, 0x3A)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    SOFT = RGBColor(0xC8, 0xD3, 0xEA)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def rect(slide, x, y, w, h, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        return shape

    def textbox(slide, x, y, w, h, text, size, color, bold=False,
                anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        run = frame.paragraphs[0].add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
        return box

    # Title slide
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, W, H, NAVY)
    textbox(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.0), data["title"],
            44, WHITE, True, MSO_ANCHOR.BOTTOM)
    rect(s, Inches(0.8), Inches(3.8), Inches(1.6), Inches(0.08), ACCENT)
    textbox(s, Inches(0.8), Inches(4.1), Inches(11), Inches(0.8),
            "Created with StudyBuddy", 20, SOFT)

    # Content slides (max 6 points per slide)
    for sec in data["sections"]:
        points = sec["points"]
        chunks = [points[i:i + 6] for i in range(0, len(points), 6)] or [[]]
        for idx, chunk in enumerate(chunks):
            s = prs.slides.add_slide(blank)
            rect(s, 0, 0, W, H, LIGHT)
            rect(s, 0, 0, W, Inches(1.3), NAVY)
            rect(s, 0, Inches(1.3), W, Inches(0.07), ACCENT)
            heading = sec["heading"] + (" (cont.)" if idx else "")
            textbox(s, Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.9),
                    heading, 30, WHITE, True, MSO_ANCHOR.MIDDLE)

            longest = max((len(p) for p in chunk), default=0)
            size = 22 if longest <= 90 else 19 if longest <= 150 else 16
            body = s.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(11.5), Inches(5.2))
            frame = body.text_frame
            frame.word_wrap = True
            for n, point in enumerate(chunk):
                para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                para.space_after = Pt(12)
                run = para.add_run()
                run.text = "\u2022  " + point
                run.font.size, run.font.color.rgb = Pt(size), TEXT

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_docx(data: dict) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(data["title"], 0)
    for sec in data["sections"]:
        doc.add_heading(sec["heading"], level=1)
        for point in sec["points"]:
            doc.add_paragraph(point, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _latin(text: str) -> str:
    """The basic PDF fonts only support Latin letters, so clean the text."""
    for a, b in {"\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"',
                 "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-"}.items():
        text = text.replace(a, b)
    text = text.encode("latin-1", "replace").decode("latin-1")
    return re.sub(r"(\S{60})(?=\S)", r"\1 ", text)  # break very long words


def build_pdf(data: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 20)
    pdf.multi_cell(0, 10, _latin(data["title"]), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for sec in data["sections"]:
        pdf.set_font("Helvetica", "B", 14)
        pdf.multi_cell(0, 8, _latin(sec["heading"]), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for point in sec["points"]:
            pdf.multi_cell(0, 6, _latin("- " + point), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
    return bytes(pdf.output())


FORMATS = {
    "PowerPoint (.pptx)": (build_pptx, "pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "PDF (.pdf)": (build_pdf, "pdf", "application/pdf"),
    "Word (.docx)": (build_docx, "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
}


def safe_name(title: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")[:40] or "StudyBuddy"


# ---------------- 3. Understand "make me a PPT on ..." typed in the chat ----------------
_MAKE = re.compile(r"\b(make|create|generate|build|prepare|produce|convert|turn|export|"
                   r"give|write|draft|need|want|download|get)\b")
_START = re.compile(r"^\s*(a |an )?(pptx?|power\s?point|pdf|presentation|slides?|docx?)\s+"
                    r"(on|about|for|of)\b")
_FORMAT_WORDS = {
    "PowerPoint (.pptx)": re.compile(r"\b(pptx?|power\s?point|slides?|slide\s?deck|presentation|deck)\b"),
    "PDF (.pdf)": re.compile(r"\bpdf\b"),
    "Word (.docx)": re.compile(r"\b(docx|ms\s?word|word\s+(file|document|doc)|"
                               r"(in|as|to|into)\s+(a\s+)?word)\b"),
}
# "...from THIS pdf" means the pdf is the source, not the file we should make
_SOURCE_BEFORE = re.compile(r"\b(this|the|attached|uploaded|my|from|of|read|summari[sz]e)\s+$")
_ANSWER_WORDS = re.compile(r"\b(above|previous|last answer|this answer|that answer|your answer|"
                           r"your last|the answer|these notes|this explanation)\b")
_FILE_WORDS = re.compile(r"\b(attach\w*|upload\w*|this (file|pdf|doc\w*|image|photo|notes)|"
                         r"these (files|notes)|the (file|pdf|document)|my (file|notes|pdf))\b")


def detect_file_request(text: str):
    """Return 'PowerPoint (.pptx)', 'PDF (.pdf)' or 'Word (.docx)' if the student
    is asking us to MAKE such a file, otherwise None."""
    low = text.lower()
    if not (_MAKE.search(low) or _START.search(low)):
        return None
    found = []
    for fmt, pat in _FORMAT_WORDS.items():
        for m in pat.finditer(low):
            before = low[max(0, m.start() - 14):m.start()]
            if _SOURCE_BEFORE.search(before):
                continue
            found.append((m.start(), fmt))
    return min(found)[1] if found else None


def plan_file_request(text: str, has_files: bool, has_answer: bool) -> dict:
    """Decide what the file should be made from and how many slides/sections."""
    low = text.lower()
    n = re.search(r"\b(\d{1,2})\s*(?:slides?|pages?|sections?)\b", low)
    n_sections = min(12, max(3, int(n.group(1)))) if n else 6
    if has_answer and _ANSWER_WORDS.search(low):
        return {"kind": "answer", "n": n_sections}
    if _FILE_WORDS.search(low):
        return {"kind": "files" if has_files else "need_file", "n": n_sections}
    return {"kind": "topic", "n": n_sections}
